"""
extract_cv.py — Extrae un perfil estructurado de un CV .docx a JSON.
No usa Ollama, solo regex y python-docx. Rápido, se corre una vez por CV
(o cada vez que actualices el CV).

Uso:
    python extract_cv.py mi_cv.docx
"""

import sys
import json
import re
from docx import Document

# Lista de herramientas conocidas, solo para el resumen que se imprime en
# pantalla y para el listado "herramientas" del JSON. match.py NO depende
# de esta lista para decidir si tienes o no una skill: siempre revisa
# también el texto completo del CV, así que una herramienta que falte aquí
# no te va a generar un falso "te falta" en match.py.
HERRAMIENTAS_CONOCIDAS = [
    "React", "TypeScript", "JavaScript", "Python", "Java", "C++", "C#",
    "Node.js", "Express", "Django", "Flask", "FastAPI",
    "MapLibre", "Mapbox", "Leaflet", "OpenLayers", "GIS", "QGIS", "ArcGIS",
    "ArcPy", "GDAL", "PostGIS", "GeoPandas", "Shapely", "rasterio", "xarray",
    "PostgreSQL", "MySQL", "MongoDB", "Redis", "SQLite", "DuckDB",
    "Docker", "Kubernetes", "AWS", "Azure", "GCP", "Lambda",
    "Git", "GitHub", "GitLab", "CI/CD", "Jenkins", "Airflow",
    "REST", "GraphQL", "API",
    "HTML", "CSS", "SASS", "Tailwind",
    "Linux", "WSL", "Bash", "PowerShell",
    "Pandas", "NumPy", "Polars", "SQLAlchemy",
    "TensorFlow", "PyTorch", "scikit-learn", "scikit-image", "OpenCV",
    "Plotly", "Matplotlib", "Power BI", "Power Query", "Power Apps",
    "Excel", "PowerPoint", "Word",
    "Figma", "Jira", "Confluence", "Slack",
]


def iter_paragraphs_incl_tables(doc: Document):
    """Recorre párrafos del cuerpo Y de celdas de tabla. Muchos CVs (el tuyo
    incluido) meten la tabla de skills técnicos en una tabla de Word, y
    doc.paragraphs por sí solo NO entra ahí."""
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para


def extraer_texto_docx(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in iter_paragraphs_incl_tables(doc) if p.text.strip())


def contiene_termino(texto: str, termino: str) -> bool:
    """Busca `termino` en `texto` como palabra/token completo, sin depender
    de \\b (que falla con términos que terminan en símbolo, como C++ o C#:
    \\b nunca marca límite justo después de un símbolo). En vez de eso,
    exige que los caracteres alfanuméricos alrededor del término no
    continúen la palabra."""
    patron = r'(?<![A-Za-z0-9])' + re.escape(termino) + r'(?![A-Za-z0-9])'
    return re.search(patron, texto, re.IGNORECASE) is not None


def extraer_perfil(texto: str) -> dict:
    encontradas = sorted({h for h in HERRAMIENTAS_CONOCIDAS if contiene_termino(texto, h)})

    anios = sorted(set(re.findall(r'\b(19\d{2}|20\d{2})\b', texto)), reverse=True)
    emails = re.findall(r'[\w\.-]+@[\w\.-]+\.\w+', texto)
    urls = re.findall(r'https?://[^\s]+', texto)

    return {
        "herramientas": encontradas,
        "anios_mencionados": anios,
        "emails": emails,
        "urls": urls,
        "texto_completo": texto,
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python extract_cv.py mi_cv.docx")
        sys.exit(1)

    texto = extraer_texto_docx(sys.argv[1])
    if not texto.strip():
        print(f"⚠️  No se extrajo texto de {sys.argv[1]}. ¿Es un .docx válido?")
        sys.exit(1)

    perfil = extraer_perfil(texto)

    salida = "cv_perfil.json"
    with open(salida, "w", encoding="utf-8") as f:
        json.dump(perfil, f, ensure_ascii=False, indent=2)

    print(f"✅ Perfil extraído de {sys.argv[1]}")
    print(f"📊 Herramientas detectadas ({len(perfil['herramientas'])}): "
          f"{', '.join(perfil['herramientas'])}")
    print(f"📅 Años mencionados: {', '.join(perfil['anios_mencionados'][:6])}")
    print(f"💾 Guardado en: {salida}")
    print("ℹ️  match.py también busca directo en el texto completo del CV, "
          "así que una herramienta que no esté en la lista de arriba igual "
          "se detecta si la JD la pide.")
